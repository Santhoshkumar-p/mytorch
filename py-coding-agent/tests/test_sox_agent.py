"""
Tests for SOX Automation Agent

Tests the agent's ability to:
- Load test templates from Excel
- Execute test cases in a loop
- Collect evidence and audit trail
- Generate compliance reports
"""

import asyncio
import json
import tempfile
from pathlib import Path
from datetime import datetime

import pytest

from sox_automation_agent import (
    SOXAutomationAgent,
    TestCase,
    TestStep,
    TestStatus,
    AuditTrailEntry,
)


@pytest.fixture
def temp_dir():
    """Create temporary directory for test artifacts"""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def sample_excel_template(temp_dir):
    """Create a sample Excel template for testing"""
    from create_sox_template import create_sox_template
    
    template_path = temp_dir / "test_template.xlsx"
    create_sox_template(str(template_path))
    return template_path


class TestSOXAgentInitialization:
    """Test agent initialization and setup"""
    
    def test_agent_creates_output_dir(self, temp_dir):
        """Agent should create output directory if it doesn't exist"""
        output_dir = temp_dir / "reports"
        agent = SOXAutomationAgent(str(temp_dir / "nonexistent.xlsx"), str(output_dir))
        
        assert output_dir.exists()
        assert agent.report is not None
        assert agent.test_cases == []
    
    def test_agent_report_has_unique_id(self, temp_dir):
        """Each agent should have a unique report ID"""
        agent1 = SOXAutomationAgent(str(temp_dir / "test1.xlsx"), str(temp_dir))
        agent2 = SOXAutomationAgent(str(temp_dir / "test2.xlsx"), str(temp_dir))
        
        assert agent1.report.report_id != agent2.report.report_id
        assert agent1.report.report_id.startswith("SOX-")
        assert agent2.report.report_id.startswith("SOX-")


class TestTemplateLoading:
    """Test loading test templates from Excel"""
    
    def test_load_template_creates_test_cases(self, sample_excel_template, temp_dir):
        """Loading template should create test cases"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        assert len(agent.test_cases) > 0
        assert all(isinstance(tc, TestCase) for tc in agent.test_cases)
    
    def test_load_template_creates_test_steps(self, sample_excel_template, temp_dir):
        """Loading template should attach steps to test cases"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        # At least one test case should have steps
        has_steps = any(len(tc.steps) > 0 for tc in agent.test_cases)
        assert has_steps
    
    def test_load_template_validates_file_exists(self, temp_dir):
        """Loading non-existent file should raise error"""
        agent = SOXAutomationAgent(str(temp_dir / "nonexistent.xlsx"), str(temp_dir))
        
        with pytest.raises(FileNotFoundError):
            agent.load_test_template()
    
    def test_load_template_parses_parameters(self, sample_excel_template, temp_dir):
        """Template loading should parse step parameters from JSON"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        # Find a step with parameters
        steps_with_params = [
            step for tc in agent.test_cases
            for step in tc.steps
            if step.parameters
        ]
        
        assert len(steps_with_params) > 0
        # Parameters should be dicts, not strings
        for step in steps_with_params:
            assert isinstance(step.parameters, dict)


class TestTestExecution:
    """Test execution of test cases"""
    
    @pytest.mark.asyncio
    async def test_execute_single_test_case(self, sample_excel_template, temp_dir):
        """Agent should execute a single test case"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        if not agent.test_cases:
            pytest.skip("No test cases in template")
        
        test_case = agent.test_cases[0]
        await agent.execute_test_case(test_case, dry_run=True)
        
        assert test_case.status != TestStatus.PENDING
        assert test_case.start_time != ""
        assert test_case.end_time != ""
    
    @pytest.mark.asyncio
    async def test_execute_all_tests(self, sample_excel_template, temp_dir):
        """Agent should execute all tests in sequence"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        # All test cases should have been executed
        assert agent.report.total_tests == len(agent.test_cases)
        # In dry run, all should pass
        assert agent.report.passed_tests == agent.report.total_tests
    
    @pytest.mark.asyncio
    async def test_test_step_execution_records_evidence(self, sample_excel_template, temp_dir):
        """Executing a test step should generate evidence ID"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        if not agent.test_cases or not agent.test_cases[0].steps:
            pytest.skip("No test steps in template")
        
        test_case = agent.test_cases[0]
        step = test_case.steps[0]
        
        await agent.execute_test_step(test_case, step, dry_run=True)
        
        assert step.evidence_id != ""
        assert step.evidence_id.startswith("EV-")
        assert step in test_case.steps
        assert step.evidence_id in test_case.evidence_ids


class TestAuditTrail:
    """Test audit trail generation"""
    
    @pytest.mark.asyncio
    async def test_audit_trail_records_test_start(self, sample_excel_template, temp_dir):
        """Audit trail should record test start"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        if not agent.test_cases:
            pytest.skip("No test cases")
        
        test_case = agent.test_cases[0]
        await agent.execute_test_case(test_case, dry_run=True)
        
        # Should have at least one audit entry for test start
        test_entries = [
            e for e in agent.report.audit_trail
            if e.test_id == test_case.test_id
        ]
        assert len(test_entries) > 0
    
    @pytest.mark.asyncio
    async def test_audit_trail_includes_evidence_ids(self, sample_excel_template, temp_dir):
        """Audit trail entries should include evidence IDs"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        # Entries for actual test steps should have evidence IDs
        step_entries = [
            e for e in agent.report.audit_trail
            if e.step_id != ""
        ]
        
        assert len(step_entries) > 0
        assert all(e.evidence_id != "" for e in step_entries)


class TestReportGeneration:
    """Test report generation and export"""
    
    @pytest.mark.asyncio
    async def test_finalize_report_calculates_totals(self, sample_excel_template, temp_dir):
        """Finalizing report should calculate all totals"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        assert agent.report.total_tests > 0
        assert agent.report.total_tests == len(agent.test_cases)
        assert agent.report.passed_tests + agent.report.failed_tests >= 0
    
    @pytest.mark.asyncio
    async def test_control_results_aggregated(self, sample_excel_template, temp_dir):
        """Report should aggregate results by control objective"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        assert len(agent.report.control_results) > 0
        # Each control should have a pass/fail status
        for control, status in agent.report.control_results.items():
            assert isinstance(control, str)
            assert isinstance(status, bool)
    
    @pytest.mark.asyncio
    async def test_export_json_report(self, sample_excel_template, temp_dir):
        """Agent should export report as JSON"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        json_file = agent.export_json_report("test_report.json")
        
        assert Path(json_file).exists()
        
        with open(json_file, 'r') as f:
            report_data = json.load(f)
        
        assert report_data['report_id'] == agent.report.report_id
        assert 'total_tests' in report_data
        assert 'test_cases' in report_data
        assert 'audit_trail' in report_data
    
    @pytest.mark.asyncio
    async def test_export_excel_report(self, sample_excel_template, temp_dir):
        """Agent should export report as Excel"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        excel_file = agent.export_excel_report("test_report.xlsx")
        
        assert Path(excel_file).exists()
        
        # Verify Excel structure
        from openpyxl import load_workbook
        wb = load_workbook(excel_file)
        
        expected_sheets = ["Summary", "Controls", "Test Cases", "Audit Trail"]
        for sheet_name in expected_sheets:
            assert sheet_name in wb.sheetnames


class TestActionHandlers:
    """Test individual action handlers"""
    
    @pytest.mark.asyncio
    async def test_verify_access_action(self, temp_dir):
        """Test verify_access action handler"""
        agent = SOXAutomationAgent(str(temp_dir / "test.xlsx"), str(temp_dir))
        
        step = TestStep(
            step_id="1",
            description="Test access",
            action="verify_access",
            expected_result="Access verified",
            parameters={"user": "test_user", "resource": "system"}
        )
        
        result = await agent._verify_access(step, dry_run=False)
        
        assert "test_user" in result
        assert "system" in result
    
    @pytest.mark.asyncio
    async def test_check_audit_trail_action(self, temp_dir):
        """Test check_audit_trail action handler"""
        agent = SOXAutomationAgent(str(temp_dir / "test.xlsx"), str(temp_dir))
        
        step = TestStep(
            step_id="1",
            description="Check audit trail",
            action="check_audit_trail",
            expected_result="Entries found",
            parameters={"action": "login", "expected_count": 1}
        )
        
        result = await agent._check_audit_trail(step, dry_run=False)
        
        assert "login" in result
    
    @pytest.mark.asyncio
    async def test_validate_changes_action(self, temp_dir):
        """Test validate_changes action handler"""
        agent = SOXAutomationAgent(str(temp_dir / "test.xlsx"), str(temp_dir))
        
        step = TestStep(
            step_id="1",
            description="Validate changes",
            action="validate_changes",
            expected_result="Changes logged",
            parameters={"resource": "permission", "change_type": "grant"}
        )
        
        result = await agent._validate_changes(step, dry_run=False)
        
        assert "permission" in result
        assert "grant" in result


class TestDryRunMode:
    """Test dry-run execution mode"""
    
    @pytest.mark.asyncio
    async def test_dry_run_does_not_execute_actions(self, sample_excel_template, temp_dir):
        """Dry-run should simulate actions without actually executing them"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        # All tests should pass in dry-run
        assert agent.report.failed_tests == 0
        assert agent.report.passed_tests == agent.report.total_tests
    
    @pytest.mark.asyncio
    async def test_dry_run_still_generates_reports(self, sample_excel_template, temp_dir):
        """Dry-run should still generate audit trail and reports"""
        agent = SOXAutomationAgent(str(sample_excel_template), str(temp_dir))
        agent.load_test_template()
        
        await agent.execute_tests(dry_run=True)
        
        assert len(agent.report.audit_trail) > 0
        assert len(agent.report.test_cases) > 0


class TestDataStructures:
    """Test data structure behavior"""
    
    def test_test_case_default_status(self):
        """Test case should start in PENDING status"""
        tc = TestCase(
            test_id="T1",
            title="Test",
            description="Test case",
            control_objective="Control",
            test_type="Design",
            frequency="Quarterly",
        )
        
        assert tc.status == TestStatus.PENDING
        assert tc.start_time == ""
        assert tc.passed_steps == 0
        assert tc.failed_steps == 0
    
    def test_test_step_default_status(self):
        """Test step should start in PENDING status"""
        step = TestStep(
            step_id="S1",
            description="Step",
            action="verify_access",
            expected_result="Pass",
        )
        
        assert step.status == TestStatus.PENDING
        assert step.timestamp == ""
        assert step.evidence_id == ""
        assert step.actual_result == ""
    
    def test_audit_trail_entry_creation(self):
        """Audit trail entries should be created correctly"""
        entry = AuditTrailEntry(
            timestamp="2026-04-25T10:00:00",
            action="test_start",
            actor="agent",
            test_id="T1",
            result="Started"
        )
        
        assert entry.timestamp == "2026-04-25T10:00:00"
        assert entry.action == "test_start"
        assert entry.actor == "agent"
        assert entry.test_id == "T1"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
